from pathlib import Path
from typing import Dict

import pandas as pd
from pptx import Presentation
from pptx.util import Inches


def export_cleaned_excel(standardized_df: pd.DataFrame, notes_df: pd.DataFrame, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        standardized_df.to_excel(writer, sheet_name="standardized_data", index=False)
        notes_df.to_excel(writer, sheet_name="data_quality_notes", index=False)
    return output_path


def export_summary_text(text: str, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(text, encoding="utf-8")
    return output_path


def export_powerpoint(chart_paths: Dict[str, str], summary_text: str, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()

    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "ATLAS Media Chart Builder"
    slide.placeholders[1].text = "Auto-generated MVP deck"

    for chart_title, chart_path in chart_paths.items():
        content_layout = prs.slide_layouts[5]
        s = prs.slides.add_slide(content_layout)
        s.shapes.title.text = chart_title
        s.shapes.add_picture(chart_path, Inches(0.8), Inches(1.5), width=Inches(8.0), height=Inches(4.5))
        tx = s.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(8.5), Inches(0.8))
        tx.text_frame.text = summary_text.splitlines()[0] if summary_text else "Summary insight"

    prs.save(output_path)
    return output_path
